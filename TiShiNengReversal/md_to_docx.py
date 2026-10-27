import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

# 字体设置
FONT_BODY = "微软雅黑"
FONT_CODE = "Consolas"
FONT_SIZE_BODY = Pt(10.5)
FONT_SIZE_H1 = Pt(22)
FONT_SIZE_H2 = Pt(16)
FONT_SIZE_H3 = Pt(13)
FONT_SIZE_CODE = Pt(9)

# 代码块颜色映射（语言 -> 背景色/前景色）
CODE_COLORS = {
    "python": {"bg": "F0F0F0", "fg": "333333"},
    "bash": {"bg": "F0F0F0", "fg": "333333"},
    "powershell": {"bg": "F0F0F0", "fg": "333333"},
    "json": {"bg": "F0F0F0", "fg": "333333"},
    "xml": {"bg": "F0F0F0", "fg": "333333"},
    "java": {"bg": "F0F0F0", "fg": "333333"},
    "javascript": {"bg": "F0F0F0", "fg": "333333"},
    "": {"bg": "F5F5F5", "fg": "333333"},  # 默认
}

# 代码块语法高亮色（关键字/字符串/注释等）
SYNTAX_COLORS = {
    "keyword": RGBColor(0, 0, 0xFF),    # 蓝色
    "string": RGBColor(0xA3, 0x15, 0x15),  # 红色
    "comment": RGBColor(0x00, 0x80, 0x00),  # 绿色
    "default": RGBColor(0x33, 0x33, 0x33),  # 深灰
}


def set_run_font(run, font_name=FONT_BODY, font_size=FONT_SIZE_BODY, bold=False, italic=False, color=None):
    """设置 run 的字体属性"""
    run.font.name = font_name
    # 设置东亚字体
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")}/>')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = font_size
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color


def add_code_block(doc, code_text, language=""):
    """添加代码块（带背景色的表格模拟）"""
    # 用单单元格表格模拟代码块背景
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    cell = table.cell(0, 0)
    # 设置单元格背景色
    bg_color = CODE_COLORS.get(language, CODE_COLORS[""])["bg"]
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
    cell._element.get_or_add_tcPr().append(shading)

    # 设置单元格边距
    tc = cell._element
    tcPr = tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="100" w:type="dxa"/>'
        f'  <w:left w:w="150" w:type="dxa"/>'
        f'  <w:bottom w:w="100" w:type="dxa"/>'
        f'  <w:right w:w="150" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

    # 清除默认段落
    cell.paragraphs[0].clear()

    # 逐行添加代码
    lines = code_text.split('\n')
    for i, line in enumerate(lines):
        if i == 0:
            p = cell.paragraphs[0]
        else:
            p = cell.add_paragraph()

        # 设置段落格式（紧凑行距）
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = Pt(14)

        run = p.add_run(line)
        set_run_font(run, FONT_CODE, FONT_SIZE_CODE, color=SYNTAX_COLORS["default"])

    # 设置表格边框为无
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else parse_xml(f'<w:tblPr {nsdecls("w")}/>')
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'  <w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'  <w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'  <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'  <w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'  <w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def add_md_table(doc, header_row, data_rows):
    """添加 md 表格"""
    cols = len(header_row)
    table = doc.add_table(rows=1 + len(data_rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # 设置表格样式（带边框）
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else parse_xml(f'<w:tblPr {nsdecls("w")}/>')
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

    # 表头（加粗，浅灰背景）
    for j, cell_text in enumerate(header_row):
        cell = table.cell(0, j)
        # 清空单元格内容
        for p in cell.paragraphs:
            for run in p.runs:
                run.clear()
        p = cell.paragraphs[0]
        run = p.add_run(cell_text.strip())
        set_run_font(run, FONT_BODY, FONT_SIZE_BODY, bold=True)
        # 表头背景色
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="E8E8E8"/>')
        cell._element.get_or_add_tcPr().append(shading)

    # 数据行
    for i, row in enumerate(data_rows):
        for j, cell_text in enumerate(row):
            if j < cols:
                cell = table.cell(i + 1, j)
                cell.text = ""
                p = cell.paragraphs[0]
                # 解析单元格内的格式（粗体、行内代码）
                add_formatted_text(p, cell_text.strip())


def add_formatted_text(paragraph, text, base_bold=False, base_italic=False):
    """解析 md 行内格式并添加到段落"""
    # 处理粗体、斜体、行内代码
    # 模式：**粗体**、*斜体*、`代码`
    parts = re.split(r'(\*\*.*?\*\*|`[^`]+`|\*[^*]+\*)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            inner = part[2:-2]
            run = paragraph.add_run(inner)
            set_run_font(run, FONT_BODY, FONT_SIZE_BODY, bold=True, italic=base_italic)
        elif part.startswith('`') and part.endswith('`'):
            inner = part[1:-1]
            run = paragraph.add_run(inner)
            set_run_font(run, FONT_CODE, FONT_SIZE_CODE, color=RGBColor(0xC7, 0x25, 0x4E))
        elif part.startswith('*') and part.endswith('*') and not part.startswith('**'):
            inner = part[1:-1]
            run = paragraph.add_run(inner)
            set_run_font(run, FONT_BODY, FONT_SIZE_BODY, bold=base_bold, italic=True)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, FONT_BODY, FONT_SIZE_BODY, bold=base_bold, italic=base_italic)


def parse_md_table(lines, start_idx):
    """解析 md 表格，返回 (header, data_rows, end_idx)"""
    header_line = lines[start_idx].strip()
    cells = [c.strip() for c in header_line.split('|')[1:-1]]

    # 跳过分隔行
    sep_idx = start_idx + 1
    data_rows = []
    idx = sep_idx + 1
    while idx < len(lines):
        line = lines[idx].strip()
        if not line.startswith('|'):
            break
        row = [c.strip() for c in line.split('|')[1:-1]]
        data_rows.append(row)
        idx += 1

    return cells, data_rows, idx


def md_to_docx(md_path, docx_path):
    """将 md 文件转换为 docx"""
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.split('\n')
    doc = Document()

    # 设置默认字体
    style = doc.styles['Normal']
    style.font.name = FONT_BODY
    style.font.size = FONT_SIZE_BODY
    style._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_BODY)

    # 设置页边距
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # 空行
        if not stripped:
            i += 1
            continue

        # 代码块
        if stripped.startswith('```'):
            language = stripped[3:].strip()
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            i += 1  # 跳过结束的 ```
            add_code_block(doc, '\n'.join(code_lines), language)
            continue

        # 表格
        if stripped.startswith('|') and i + 1 < len(lines) and '---' in lines[i + 1]:
            header, data_rows, end_idx = parse_md_table(lines, i)
            add_md_table(doc, header, data_rows)
            i = end_idx
            continue

        # 一级标题 (# )
        if stripped.startswith('# ') and not stripped.startswith('## '):
            text = stripped[2:].strip()
            p = doc.add_heading(level=1)
            run = p.add_run(text)
            set_run_font(run, FONT_BODY, FONT_SIZE_H1, bold=True)
            i += 1
            continue

        # 二级标题 (## )
        if stripped.startswith('## ') and not stripped.startswith('### '):
            text = stripped[3:].strip()
            p = doc.add_heading(level=2)
            run = p.add_run(text)
            set_run_font(run, FONT_BODY, FONT_SIZE_H2, bold=True)
            i += 1
            continue

        # 三级标题 (### )
        if stripped.startswith('### '):
            text = stripped[4:].strip()
            p = doc.add_heading(level=3)
            run = p.add_run(text)
            set_run_font(run, FONT_BODY, FONT_SIZE_H3, bold=True)
            i += 1
            continue

        # 引用块 (>)
        if stripped.startswith('>'):
            text = stripped[1:].strip()
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Cm(1)
            pf.space_before = Pt(2)
            pf.space_after = Pt(2)
            # 引用块左侧边框
            pPr = p._element.get_or_add_pPr()
            pBdr = parse_xml(
                f'<w:pBdr {nsdecls("w")}>'
                f'  <w:left w:val="single" w:sz="12" w:space="4" w:color="CCCCCC"/>'
                f'</w:pBdr>'
            )
            pPr.append(pBdr)
            add_formatted_text(p, text)
            i += 1
            continue

        # 无序列表 (- )
        if stripped.startswith('- '):
            text = stripped[2:].strip()
            # 检查是否有缩进（子列表）
            indent = len(line) - len(line.lstrip())
            p = doc.add_paragraph(style='List Bullet')
            if indent > 0:
                p.paragraph_format.left_indent = Cm(1.5 + (indent // 2) * 0.5)
            add_formatted_text(p, text)
            i += 1
            continue

        # 有序列表 (1. 2. 等)
        if re.match(r'^\d+\.\s', stripped):
            text = re.sub(r'^\d+\.\s', '', stripped)
            p = doc.add_paragraph(style='List Number')
            add_formatted_text(p, text)
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        add_formatted_text(p, stripped)
        i += 1

    doc.save(docx_path)
    print(f"docx 已生成: {docx_path}")


if __name__ == "__main__":
    md_path = r"D:\ClaudeCode\tishineng_docx\体适能_新增不合规点与法条标注补充_替换后.md"
    docx_path = r"D:\ClaudeCode\tishineng_docx\体适能_新增不合规点与法条标注补充.docx"
    md_to_docx(md_path, docx_path)
